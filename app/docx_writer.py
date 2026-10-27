"""内评简评报告 docx 生成器（与模板 v1.1 一一对应）。

只依赖 python-docx。输入一个符合 01_模板/内评简评报告通用模板_v1.1.json 的
数据结构，输出 .docx，版式按参考件（人工撰写的同类报告）对齐：

- 中文楷体_GB2312、五号(10.5pt)、行距1.15、首行缩进2字符、段前段后0
- 标题行 / 正面： / 关注： / 1、基本情况 / 2、点评 / 汇总句 / 表题 加粗
- 段落之间插入一个空段作为分块间隔（参考件即如此，否则整页挤成一块）
- 表格：外框上下 1.5pt、内部横竖线 0.5pt、左右不封口；表头加粗居中，
  数值列右对齐、首列左对齐；行高 0.4cm 起、单元格垂直居中、禁止跨页断行
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT_CN = "楷体_GB2312"
FONT_SIZE_PT = 10.5
TEXT_WIDTH_CM = 14.66          # A4 21cm - 左右各 3.17cm
ROW_HEIGHT_TWIPS = 227         # 参考件行高 0.4cm（atLeast）

_TBLPR_TAIL = ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
               "w:tblCaption", "w:tblDescription", "w:tblPrChange")


def _end_period(text: str) -> str:
    """摘要条目统一以句号收尾，避免生成稿出现无标点短句。"""
    t = (text or "").strip()
    return t if not t or t[-1] in "。！？；" else t + "。"


def _set_run_font(run, bold: bool = False, size_pt: float = FONT_SIZE_PT):
    run.font.name = FONT_CN
    run.font.size = Pt(size_pt)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), FONT_CN)
    rfonts.set(qn("w:ascii"), FONT_CN)
    rfonts.set(qn("w:hAnsi"), FONT_CN)


def _indent_two_chars(paragraph):
    """真正的『首行缩进 2 字符』（firstLineChars=200），与 Word 行为一致。"""
    ppr = paragraph._p.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    ind.set(qn("w:firstLineChars"), "200")
    ind.set(qn("w:firstLine"), str(int(FONT_SIZE_PT * 2 * 20)))


def _fmt_paragraph(paragraph, align=None):
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.15
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    # 正文默认两端对齐：中文段落右边缘必须齐，否则满行参差。
    # 标题、表格数字等由调用方显式传 align 覆盖，不受影响。
    paragraph.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    _indent_two_chars(paragraph)


def _add_para(doc, text: str, bold: bool = False, align=None, indent: bool = True):
    p = doc.add_paragraph()
    _fmt_paragraph(p, align)
    if not indent:
        ppr = p._p.get_or_add_pPr()
        for ind in ppr.findall(qn("w:ind")):
            ppr.remove(ind)
    r = p.add_run(text)
    _set_run_font(r, bold=bold)
    return p


def _add_blank(doc):
    """块之间的空行。参考件用空段落分隔，不靠段前段后间距。"""
    p = doc.add_paragraph()
    _fmt_paragraph(p)
    return p


def _add_lead_para(doc, lead: str, body: str, bold_lead: bool = True, align=None):
    """段首加粗汇总句 + 常规正文。"""
    p = doc.add_paragraph()
    _fmt_paragraph(p, align)
    if lead:
        r1 = p.add_run(lead)
        _set_run_font(r1, bold=bold_lead)
    if body:
        r2 = p.add_run(body)
        _set_run_font(r2, bold=False)
    return p


def _cell_text(cell, text: str, bold=False, align=WD_ALIGN_PARAGRAPH.RIGHT):
    cell.text = ""
    p = cell.paragraphs[0]
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.15
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    p.alignment = align
    r = p.add_run(str(text))
    _set_run_font(r, bold=bold)
    # 单元格垂直居中，与参考件一致
    tcPr = cell._tc.get_or_add_tcPr()
    valign = tcPr.find(qn("w:vAlign"))
    if valign is None:
        valign = OxmlElement("w:vAlign")
        tcPr.append(valign)
    valign.set(qn("w:val"), "center")


def _tbl_borders(tblPr):
    """外框上下 1.5pt、左右不画；内部横竖线 0.5pt。与参考件一致。"""
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge, val, sz in (("top", "single", 12), ("left", "none", 0),
                          ("bottom", "single", 12), ("right", "none", 0),
                          ("insideH", "single", 4), ("insideV", "single", 4)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.insert_element_before(borders, *_TBLPR_TAIL)


def _tbl_cell_margins(tblPr):
    for old in tblPr.findall(qn("w:tblCellMar")):
        tblPr.remove(old)
    mar = OxmlElement("w:tblCellMar")
    for side, w in (("top", 0), ("left", 108), ("bottom", 0), ("right", 108)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(w))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.insert_element_before(mar, "w:tblLook", "w:tblCaption",
                                "w:tblDescription", "w:tblPrChange")


def _tbl_width(tblPr, cm: float):
    el = tblPr.find(qn("w:tblW"))
    if el is None:
        el = OxmlElement("w:tblW")
        tblPr.insert_element_before(el, "w:jc", "w:tblCellSpacing", "w:tblInd",
                                    *_TBLPR_TAIL)
    el.set(qn("w:type"), "dxa")
    el.set(qn("w:w"), str(int(round(cm * 567))))


def _row_props(row, header: bool = False):
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn("w:cantSplit")) is None:
        trPr.append(OxmlElement("w:cantSplit"))
    if header and trPr.find(qn("w:tblHeader")) is None:
        trPr.append(OxmlElement("w:tblHeader"))
    row.height = Pt(ROW_HEIGHT_TWIPS / 20.0)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST


def _add_table(doc, header, rows, unit_line=None, col_widths_cm=None, bold_rows=()):
    if unit_line:
        up = _add_para(doc, unit_line, align=WD_ALIGN_PARAGRAPH.RIGHT, indent=False)
        up.paragraph_format.keep_with_next = True  # 单位行不单独留在页尾
    ncols = len(header) if header else max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    if header:
        cells = table.add_row().cells
        for i, h in enumerate(header):
            _cell_text(cells[i], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i in range(ncols):
            val = row[i] if i < len(row) else ""
            # 长文本（如项目名、板块名）左对齐，纯数值右对齐 → 数值列才能看齐
            is_num = bool(val) and all(ch.isdigit() or ch in ".,%-—()（） " for ch in str(val))
            align = WD_ALIGN_PARAGRAPH.RIGHT if (i > 0 and is_num) else WD_ALIGN_PARAGRAPH.LEFT
            _cell_text(cells[i], val, bold=(idx in bold_rows), align=align)

    widths = list(col_widths_cm or [])
    if len(widths) != ncols:
        first = 6.6 if ncols <= 4 else 5.0
        rest = max(TEXT_WIDTH_CM - first, 4.0) / max(ncols - 1, 1)
        widths = [first] + [round(rest, 2)] * max(ncols - 1, 0)
    total = sum(widths)
    if total > TEXT_WIDTH_CM:  # 按版心等比压缩，避免表格超出页边距
        widths = [round(w * TEXT_WIDTH_CM / total, 2) for w in widths]

    tblPr = table._tbl.tblPr
    _tbl_borders(tblPr)
    _tbl_cell_margins(tblPr)
    _tbl_width(tblPr, sum(widths))
    table.autofit = False

    for i, w in enumerate(widths):
        table.columns[i].width = Cm(w)
    for idx, row in enumerate(table.rows):
        _row_props(row, header=(idx == 0 and bool(header)))
        for i, w in enumerate(widths):
            if i < len(row.cells):
                row.cells[i].width = Cm(w)
    return table


def _fin_col_widths(n_periods: int) -> list[float]:
    """财务表列宽照参考件：项目列 6.61，最新期 3.61（「2026年3月末」一行放得下），
    更早的整年列 2.4；超过 4 列时整体压缩，保证不出版心。"""
    n = max(n_periods + 1, 2)
    if n == 2:
        return [6.61, 7.5]
    if n == 3:
        return [6.61, 3.9, 3.9]
    if n == 4:
        return [6.61, 3.61, 2.4, 2.4]
    widths = [5.6, 3.4] + [2.0] * (n - 2)
    total = sum(widths)
    return [round(w * TEXT_WIDTH_CM / total, 2) for w in widths]


def render_report(data: dict[str, Any], out_path: str | Path) -> Path:
    doc = Document()
    # 页面：A4，左右 3.17cm，上下 2.54cm
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(3.17)
    sec.top_margin = sec.bottom_margin = Cm(2.54)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_CN
    normal.font.size = Pt(FONT_SIZE_PT)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)

    # ── 版式：块与块之间一个空行，块内（加粗句 + 它的正文）不留空 ──
    #
    # 空行的位置只有一处定义，就是这里的 _gap()：每写完一个「块」调一次。
    # 段内不留空，是因为加粗总起句和它的正文本来就在同一个段落里（_add_lead_para），
    # 一旦在两者之间插空段，读起来就像话没说完就断了。
    gap_on = (data.get("layout") or {}).get("para_gap") is not False

    def _gap():
        if gap_on:
            _add_blank(doc)

    company = data.get("company_full_name", "")
    code = (data.get("internal_rating_code") or "").strip()
    _add_para(doc, f"{company}  主体：【{code or '待填'}】", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
    _gap()

    _add_para(doc, "正面：", bold=True)
    for i, item in enumerate(data.get("positive") or [], 1):
        _add_para(doc, f"{i}）{_end_period(item)}")
    _gap()

    _add_para(doc, "关注：", bold=True)
    for i, item in enumerate(data.get("concern") or [], 1):
        _add_para(doc, f"{i}）{_end_period(item)}")
    _gap()

    _add_para(doc, "1、基本情况", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
    _add_para(doc, data.get("basic_info", ""))
    _gap()

    _add_para(doc, "2、点评", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
    c = data.get("commentary") or {}

    ov = c.get("overview") or {}
    if ov:
        _add_lead_para(doc, ov.get("lead", ""), ov.get("body", ""))

    seg = c.get("segment_table")
    if seg:
        # 总述段与业务构成表是两个块，中间要有空行（原来这里漏了，
        # 表就顶在总述段的最后一行上）
        _gap()
        _add_table(
            doc,
            seg.get("header") or ["业务板块", "营业收入", "收入占比", "毛利率"],
            seg.get("rows") or [],
            unit_line=seg.get("unit_line", "单位：亿元"),
            col_widths_cm=seg.get("col_widths_cm") or [4.91, 3.91, 3.1, 3.1],
            bold_rows=[len(seg.get("rows") or []) - 1] if seg.get("last_row_bold", True) else [],
        )
    _gap()

    from generator import segment_heading  # 与网页预览、校验共用同一套编号口径

    for i, s in enumerate(c.get("segments") or [], 1):
        name = s.get("name", "")
        body = s.get("body") or s.get("text") or ""
        _add_para(doc, f"{segment_heading(i, name)}{body}")
        _gap()          # 一个板块就是一段，段与段之间同样一个空行
    if not c.get("segments"):
        pass

    at = c.get("asset") or {}
    if at:
        _add_lead_para(doc, at.get("lead", ""), at.get("body", ""))
        _gap()

    ft = c.get("financial_table")
    if ft:
        cap = _add_para(doc, ft.get("caption", "公司主要财务数据和财务指标（合并口径）"),
                        bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        cap.paragraph_format.keep_with_next = True  # 表题与表格同页
        _add_table(
            doc,
            ["项目"] + list(ft.get("periods") or []),
            ft.get("rows") or [],
            unit_line=ft.get("unit_line", "单位：亿元"),
            col_widths_cm=ft.get("col_widths_cm") or _fin_col_widths(len(ft.get("periods") or [])),
        )
        _gap()

    db = c.get("debt") or {}
    if db:
        _add_lead_para(doc, db.get("lead", ""), db.get("body", ""))
        _gap()

    topics = c.get("special_topics") or []
    for t in topics:
        if isinstance(t, str):
            _add_para(doc, t)
        else:
            _add_lead_para(doc, t.get("lead", ""), t.get("body", ""))
        _gap()

    # 结尾多余的空段去掉：Word 末尾挂一个空行，页数会多出一页空白
    while doc.paragraphs and not doc.paragraphs[-1].text.strip() and not doc.paragraphs[-1].runs:
        p = doc.paragraphs[-1]._p
        p.getparent().remove(p)

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


def render_skeleton(template_json: str | Path, out_path: str | Path) -> Path:
    """按模板 JSON 生成空白骨架 docx（槽位以 {{}} 占位），供人工套写。"""
    tpl = json.loads(Path(template_json).read_text(encoding="utf-8"))
    data = {
        "company_full_name": "{{企业全称}}",
        "internal_rating_code": "{{内部评级代码}}",
        "positive": ["{{正面一}}", "{{正面二}}", "{{正面三}}"],
        "concern": ["{{关注一}}", "{{关注二}}", "{{关注三}}"],
        "basic_info": "{{基本情况：成立时间与前身 → 主营与功能定位 → 注册资本、实收资本、股东及持股比例、实际控制人 → 报告期内股权或控制权变动 → 增资到位、上市状态、外部主体评级、并表子公司数量}}",
        "commentary": {
            "overview": {
                "lead": "{{总述段加粗汇总句：主体定位 + 经营趋势 + 但书}}",
                "body": "{{外部条件按三问展开（钱从谁那里来 / 量由什么决定 / 价由谁定）→ 经营与盈利主线数据 → 收入构成概览，末句引出业务构成表}}",
            },
            "segment_table": {
                "unit_line": "单位：亿元",
                "header": ["业务板块", "营业收入", "收入占比", "毛利率"],
                "rows": [
                    ["{{板块一}}", "{{}}", "{{}}", "{{}}"],
                    ["{{板块二}}", "{{}}", "{{}}", "{{}}"],
                    ["营业总收入", "{{}}", "100.00%", "{{}}"],
                ],
            },
            "segments": [
                {"name": "{{板块一}}", "body": "{{业务模式与产出能力 → 报告期产出量（计量单位按业务性质选取）→ 价格或费率与收入 → 成本结构与毛利率变动原因 → 项目或订单储备 → 这一块受什么制约（接在前面的经营事实上写，不另起一句「该板块风险在于」）}}"},
                {"name": "{{板块二}}", "body": "{{同上}}"},
            ],
            "asset": {
                "lead": "{{资产段加粗汇总句：规模 + 结构 + 质量 + 受限}}",
                "body": "{{总资产规模 → 前五大科目逐项：科目 金额（占资产总额 ××.××%），主要为×××，能否变现／是否质押／账龄 → 前五大合计占比 → 受限资产（金额、占比、主要为哪些）}}",
            },
            "financial_table": {
                "caption": "公司主要财务数据和财务指标（合并口径）",
                "unit_line": "单位：亿元",
                "periods": ["{{最新期末}}", "{{最近完整年度}}", "{{上一年度}}"],
                "rows": [
                    ["资产总额", "{{}}", "{{}}", "{{}}"],
                    ["负债总额", "{{}}", "{{}}", "{{}}"],
                    ["所有者权益", "{{}}", "{{}}", "{{}}"],
                    ["资产负债率", "{{}}", "{{}}", "{{}}"],
                    ["营业总收入", "{{}}", "{{}}", "{{}}"],
                    ["净利润", "{{}}", "{{}}", "{{}}"],
                    ["经营活动现金流量净额", "{{}}", "{{}}", "{{}}"],
                    ["短期债务", "{{}}", "{{}}", "{{}}"],
                    ["长期债务", "{{}}", "{{}}", "{{}}"],
                    ["全部债务", "{{}}", "{{}}", "{{}}"],
                    ["总资本化比率", "{{}}", "{{}}", "{{}}"],
                    ["营业毛利率", "{{}}", "{{}}", "{{}}"],
                    ["EBITDA", "{{}}", "{{}}", "{{}}"],
                    ["货币资金/短期债务（倍）", "{{}}", "{{}}", "{{}}"],
                ],
            },
            "debt": {
                "lead": "{{债务段加粗汇总句：规模趋势 + 期限结构 + 即期压力 + 备用流动性}}",
                "body": "{{负债总额与资产负债率 → 全部债务及长短债结构 → 有息债务融资构成 → 利息保障与经营获现 → 非受限货币资金/短债 → 授信与未使用额度}}",
            },
            "special_topics": [],
        },
    }
    return render_report(data, out_path)


if __name__ == "__main__":
    import sys

    base = Path(__file__).resolve().parent.parent
    tpl = base / "01_模板" / "内评简评报告通用模板_v1.1.json"
    out = base / "01_模板" / "内评简评报告通用模板_v1.1.docx"
    if len(sys.argv) > 1 and sys.argv[1] == "skeleton":
        print(render_skeleton(tpl, out))
    else:
        print("usage: docx_writer.py skeleton")
